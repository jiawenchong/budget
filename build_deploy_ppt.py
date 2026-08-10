# -*- coding: utf-8 -*-
"""Flask + IIS 部署新手操作簡報（黑底主題，沿用專案報告風格）。"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ---- 黑底主題色 ----
BG    = RGBColor(0x10, 0x12, 0x16)
PANEL = RGBColor(0x1B, 0x1F, 0x27)
DARK  = RGBColor(0x2A, 0x2F, 0x39)
GREYB = RGBColor(0xC9, 0xCF, 0xD8)
CYAN  = RGBColor(0x35, 0xD0, 0xE6)
BLUE  = RGBColor(0x3E, 0x8F, 0xD6)
GOLD  = RGBColor(0xE7, 0xB4, 0x3C)
GREEN = RGBColor(0x4F, 0xC1, 0x7A)
RED   = RGBColor(0xE0, 0x6C, 0x6C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTE  = RGBColor(0x8A, 0x93, 0xA1)
CODE  = RGBColor(0x9C, 0xDC, 0xFE)
FONT  = "Microsoft JhengHei"
MONO  = "Consolas"

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def slide():
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    r.fill.solid(); r.fill.fore_color.rgb = BG; r.line.fill.background()
    r.shadow.inherit = False
    return s


def rect(s, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE, line=None, lw=1.0):
    sp = s.shapes.add_shape(shape, x, y, w, h)
    sp.fill.solid(); sp.fill.fore_color.rgb = color
    if line: sp.line.color.rgb = line; sp.line.width = Pt(lw)
    else: sp.line.fill.background()
    sp.shadow.inherit = False
    return sp


def txt(s, x, y, w, h, lines, size=14, color=WHITE, bold=False,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space=3, font=FONT):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_top = 0; tf.margin_bottom = 0; tf.margin_left = Pt(2); tf.margin_right = Pt(2)
    if isinstance(lines, str): lines = [(lines, size, color, bold)]
    for i, item in enumerate(lines):
        if len(item) == 4:
            t, sz, c, b = item; fn = font
        else:
            t, sz, c, b, fn = item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(space); p.space_before = Pt(0)
        r = p.add_run(); r.text = t
        r.font.size = Pt(sz); r.font.bold = b; r.font.color.rgb = c; r.font.name = fn
    return tb


def title_bar(s, kicker, title, page):
    rect(s, 0, 0, Inches(0.18), Inches(1.05), CYAN)
    txt(s, Inches(0.5), Inches(0.14), Inches(11), Inches(0.32), kicker, 12, CYAN, True)
    txt(s, Inches(0.5), Inches(0.42), Inches(11.2), Inches(0.6), title, 26, WHITE, True)
    txt(s, Inches(11.7), Inches(0.3), Inches(1.3), Inches(0.5), f"PAGE {page}", 12, MUTE,
        True, align=PP_ALIGN.RIGHT)
    rect(s, Inches(0.5), Inches(1.08), Inches(12.33), Pt(2), DARK)


def panel(s, x, y, w, h, head, color=CYAN):
    rect(s, x, y, w, h, PANEL)
    rect(s, x, y, w, Inches(0.42), DARK)
    rect(s, x, y, Pt(6), Inches(0.42), color)
    txt(s, x + Inches(0.2), y, w - Inches(0.3), Inches(0.42), head, 14, color, True,
        anchor=MSO_ANCHOR.MIDDLE)


def codebox(s, x, y, w, h, lines):
    rect(s, x, y, w, h, RGBColor(0x0C, 0x0E, 0x12), line=DARK)
    body = [(t, 11, CODE, False, MONO) for t in lines]
    txt(s, x + Inches(0.15), y + Inches(0.1), w - Inches(0.3), h - Inches(0.2),
        body, anchor=MSO_ANCHOR.MIDDLE, space=3)


def bullets(s, x, y, w, h, items, size=12.5, space=7, color=GREYB):
    txt(s, x, y, w, h, [(t, size, c if c else color, b) for (t, c, b) in items], space=space)


# =================================================================
# 1 封面
# =================================================================
s = slide()
rect(s, 0, 0, Inches(0.25), SH, CYAN)
txt(s, Inches(0.9), Inches(2.0), Inches(11.5), Inches(0.5),
    "DEPLOYMENT GUIDE ‧ 新手專用", 14, CYAN, True)
txt(s, Inches(0.9), Inches(2.55), Inches(11.6), Inches(1.4),
    "Flask 應用部署到 IIS\n操作簡報", 40, WHITE, True, space=4)
txt(s, Inches(0.9), Inches(4.4), Inches(11.5), Inches(0.5),
    "IIS + HttpPlatformHandler + Waitress ‧ 支援 AD Windows 驗證 (SSO)", 16, GREYB, False)
rect(s, Inches(0.92), Inches(5.0), Inches(4.5), Pt(2), DARK)
txt(s, Inches(0.9), Inches(5.2), Inches(11.5), Inches(1.2),
    [("一步一步教你：打開什麼、按什麼、踩到坑怎麼救", 13, MUTE, False),
     ("範例專案：budgetagent ‧ 部署路徑 D:\\Web Service\\budgetagent ‧ Port 8001", 12, MUTE, False)],
    space=6)

# =================================================================
# 2 三個角色 + 好處
# =================================================================
s = slide()
title_bar(s, "WHY ‧ 為什麼用這套", "三個角色，與架 IIS 的好處", "02")

panel(s, Inches(0.5), Inches(1.4), Inches(6.1), Inches(2.6), "三個角色（白話）")
roles = [("IIS", "餐廳大門＋接待員：接請求、做 Windows 驗證", CYAN),
         ("HttpPlatformHandler", "IIS 外掛：自動啟動 Python、當機自動重開", BLUE),
         ("Waitress", "翻譯員：把請求翻成 Python 看得懂的格式", GREEN),
         ("Flask（你的程式）", "廚師：真正做事、回傳網頁", GOLD)]
for i, (h, d, c) in enumerate(roles):
    y = Inches(2.0) + i * Inches(0.48)
    rect(s, Inches(0.7), y + Inches(0.05), Pt(7), Inches(0.34), c)
    txt(s, Inches(0.9), y, Inches(2.6), Inches(0.42), h, 12.5, WHITE, True, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, Inches(3.4), y, Inches(3.0), Inches(0.42), d, 10.5, GREYB, False, anchor=MSO_ANCHOR.MIDDLE)

panel(s, Inches(6.83), Inches(1.4), Inches(6.0), Inches(2.6), "架 IIS 的好處", GOLD)
bens = ["✓ 員工免登入，自動認出 Windows 身分（SSO）",
        "✓ 程式當掉，IIS 自動把它重開",
        "✓ 關掉所有視窗 / 登出 / 重開機，網站照跑",
        "✓ 微軟官方推薦，取代停止維護的 wfastcgi",
        "✓ 不用另外裝服務管理工具（NSSM 等）"]
txt(s, Inches(7.03), Inches(2.0), Inches(5.6), Inches(1.9),
    [(t, 12.5, GREYB, False) for t in bens], space=10)

panel(s, Inches(0.5), Inches(4.2), Inches(12.33), Inches(2.7), "請求怎麼流動")
flow = [("瀏覽器", "自動帶\nWindows 憑證", CYAN), ("IIS", "Windows 驗證\n產生 token", BLUE),
        ("HttpPlatform\nHandler", "啟動 Python\n轉發請求", GREEN),
        ("Waitress", "轉成 WSGI\n格式", GOLD), ("Flask 應用", "解析身分\n執行業務", RGBColor(0xC7,0x7D,0xE0))]
bw = Inches(2.15)
for i, (h, d, c) in enumerate(flow):
    x = Inches(0.75) + i * Inches(2.45); y = Inches(5.0)
    rect(s, x, y, bw, Inches(1.4), BG, line=DARK)
    rect(s, x, y, bw, Pt(6), c)
    txt(s, x, y + Inches(0.18), bw, Inches(0.55), h, 13, WHITE, True, align=PP_ALIGN.CENTER)
    txt(s, x, y + Inches(0.78), bw, Inches(0.55), d, 10, MUTE, False, align=PP_ALIGN.CENTER, space=0)
    if i < 4:
        txt(s, x + bw - Inches(0.1), y + Inches(0.45), Inches(0.4), Inches(0.5),
            "▶", 14, CYAN, True, align=PP_ALIGN.CENTER)

# =================================================================
# 3 核心觀念：兩台電腦
# =================================================================
s = slide()
title_bar(s, "CORE ‧ 最重要觀念", "有「兩台電腦」，別搞混", "03")

panel(s, Inches(0.5), Inches(1.4), Inches(6.0), Inches(3.4), "🌐 有網路的電腦", GREEN)
txt(s, Inches(0.7), Inches(2.0), Inches(5.6), Inches(2.7),
    [("唯一任務：下載打包套件（產生 wheels）", 13.5, WHITE, True),
     ("", 6, MUTE, False),
     ("指令開頭：pip download", 12, GREEN, False),
     ("把套件從網路抓下來，存成一包 .whl 檔，", 12, GREYB, False),
     ("再用 USB 搬到部署機。", 12, GREYB, False),
     ("", 6, MUTE, False),
     ("公司內網要過 Proxy → 指令加 --proxy", 11.5, GOLD, False)], space=6)

panel(s, Inches(6.83), Inches(1.4), Inches(6.0), Inches(3.4), "💻 部署機", CYAN)
txt(s, Inches(7.03), Inches(2.0), Inches(5.6), Inches(2.7),
    [("其他全部都在這裡做：", 13.5, WHITE, True),
     ("", 6, MUTE, False),
     ("‧ 裝 IIS / HttpPlatformHandler / Python", 12, GREYB, False),
     ("‧ 建 venv、離線安裝套件", 12, GREYB, False),
     ("‧ 寫 wsgi.py / web.config", 12, GREYB, False),
     ("‧ 建 IIS 網站、設驗證與權限", 12, GREYB, False),
     ("‧ 上線測試", 12, GREYB, False)], space=6)

rect(s, Inches(0.5), Inches(5.1), Inches(12.33), Inches(1.7), PANEL)
rect(s, Inches(0.5), Inches(5.1), Pt(8), Inches(1.7), GOLD)
txt(s, Inches(0.8), Inches(5.25), Inches(12), Inches(0.5),
    "口訣", 15, GOLD, True)
txt(s, Inches(0.8), Inches(5.7), Inches(11.8), Inches(1.0),
    "只有「下載 wheels」要去有網路的電腦，剩下全部都在部署機。", 20, WHITE, True,
    anchor=MSO_ANCHOR.MIDDLE)

# =================================================================
# 4 事前準備
# =================================================================
s = slide()
title_bar(s, "PREP ‧ 事前準備", "開始前先把這些備齊", "04")

panel(s, Inches(0.5), Inches(1.4), Inches(6.1), Inches(2.5), "跟 IT / AD 管理員要", BLUE)
txt(s, Inches(0.7), Inches(2.0), Inches(5.7), Inches(1.9),
    [("‧ AD 短名 / 完整網域 / BaseDN", 12.5, GREYB, False),
     ("  例：KH ‧ kh.asegroup.com ‧ DC=kh,DC=asegroup,DC=com", 10.5, MUTE, False),
     ("‧ 登入格式（網域\\工號 或 UPN）", 12.5, GREYB, False),
     ("‧ ★ 對外上網的 Proxy 位址 + Port", 12.5, GOLD, False),
     ("  （下載 wheels 會用到，內網幾乎都有）", 10.5, MUTE, False)], space=7)

panel(s, Inches(6.83), Inches(1.4), Inches(6.0), Inches(2.5), "要下載的三個東西", CYAN)
txt(s, Inches(7.03), Inches(2.0), Inches(5.6), Inches(1.9),
    [("① Python 安裝檔（python.org，64-bit）", 12.5, GREYB, False),
     ("② HttpPlatformHandler v2.0 MSI（iis.net）", 12.5, GREYB, False),
     ("    IIS 不內建，要單獨裝", 10.5, MUTE, False),
     ("③ 你的專案程式碼（含 requirements.txt）", 12.5, GREYB, False)], space=9)

panel(s, Inches(0.5), Inches(4.1), Inches(12.33), Inches(2.7), "六大步驟總覽")
steps = [("1", "環境安裝", "IIS+模組+Python"), ("2", "venv+套件", "wheels 離線裝"),
         ("3", "設定檔", "wsgi.py+web.config"), ("4", "IIS 網站", "AppPool+驗證"),
         ("5", "權限", "icacls 授權"), ("6", "上線測試", "瀏覽器驗證")]
sw2 = Inches(1.92)
for i, (n, h, d) in enumerate(steps):
    x = Inches(0.7) + i * Inches(2.0); y = Inches(4.7)
    rect(s, x, y, sw2, Inches(1.5), BG, line=DARK)
    rect(s, x, y, sw2, Pt(6), CYAN)
    txt(s, x, y + Inches(0.2), sw2, Inches(0.5), f"STEP {n}", 12, CYAN, True, align=PP_ALIGN.CENTER)
    txt(s, x, y + Inches(0.62), sw2, Inches(0.45), h, 14, WHITE, True, align=PP_ALIGN.CENTER)
    txt(s, x, y + Inches(1.05), sw2, Inches(0.4), d, 10, MUTE, False, align=PP_ALIGN.CENTER)

# =================================================================
# 5 Step 1
# =================================================================
s = slide()
title_bar(s, "STEP 1 ‧ 部署機", "安裝 IIS + 模組 + Python", "05")

panel(s, Inches(0.5), Inches(1.4), Inches(6.1), Inches(2.4), "① 裝 IIS（含 Windows 驗證+CGI）", CYAN)
codebox(s, Inches(0.7), Inches(2.0), Inches(5.7), Inches(0.6),
        ["Install-WindowsFeature -Name Web-Server,",
         "  Web-Windows-Auth,Web-CGI -IncludeManagementTools"])
panel_y = Inches(2.85)
txt(s, Inches(0.7), Inches(2.85), Inches(5.7), Inches(0.9),
    [("② 裝 HttpPlatformHandler（點兩下 MSI）後驗證：", 12, WHITE, True)], space=4)
codebox(s, Inches(0.7), Inches(3.2), Inches(5.7), Inches(0.45),
        ['Get-WebGlobalModule -Name "httpPlatformHandler"'])

panel(s, Inches(6.83), Inches(1.4), Inches(6.0), Inches(2.4), "③ 裝 Python（務必勾選）", GOLD)
txt(s, Inches(7.03), Inches(2.0), Inches(5.6), Inches(1.7),
    [("☑ Install for all users  （必勾！）", 13, WHITE, True),
     ("    不勾 → 裝到 AppData → 整套失敗", 10.5, MUTE, False),
     ("☑ Add Python to PATH", 13, WHITE, True),
     ("", 6, MUTE, False),
     ("驗證裝對位置：where.exe python", 11.5, GREEN, False),
     ("要在 Program Files，不在 AppData", 10.5, MUTE, False)], space=6)

rect(s, Inches(0.5), Inches(4.0), Inches(12.33), Inches(2.8), PANEL)
rect(s, Inches(0.5), Inches(4.0), Pt(8), Inches(2.8), RED)
txt(s, Inches(0.8), Inches(4.15), Inches(12), Inches(0.4), "⚠ 關鍵檢查點", 15, RED, True)
txt(s, Inches(0.8), Inches(4.65), Inches(12), Inches(2.0),
    [("‧ HttpPlatformHandler 驗證指令「有輸出」才算裝成功。沒輸出 → 重裝 MSI。", 13, GREYB, False),
     ("   （這步沒做好，最後上線會出現 500.19 錯誤，超難找）", 11, MUTE, False),
     ("‧ Python 一定要勾「Install for all users」，否則 where.exe python 會指到 AppData。", 13, GREYB, False),
     ("‧ 安裝完記得重啟，IIS 功能才會生效。", 13, GREYB, False)], space=12)

# =================================================================
# 6 Step 2
# =================================================================
s = slide()
title_bar(s, "STEP 2 ‧ 套件安裝", "wheels 打包 + venv 離線安裝", "06")

panel(s, Inches(0.5), Inches(1.4), Inches(6.1), Inches(2.9), "在「有網路電腦」打包 wheels", GREEN)
codebox(s, Inches(0.7), Inches(2.0), Inches(5.7), Inches(1.0),
        ["pip download -r requirements.txt -d wheels",
         "  --platform win_amd64 --python-version 3.14",
         '  --only-binary=:all: --proxy "http://..."'])
txt(s, Inches(0.7), Inches(3.1), Inches(5.7), Inches(1.1),
    [("‧ --python-version 對齊「部署機」版本", 12, GREYB, False),
     ("‧ 407 錯誤 → 加 --proxy（位址問 IT）", 12, GREYB, False),
     ("‧ requirements.txt 別漏 waitress / pywin32", 12, GOLD, False)], space=6)

panel(s, Inches(6.83), Inches(1.4), Inches(6.0), Inches(2.9), "搬到部署機，建 venv 離線裝", CYAN)
codebox(s, Inches(7.03), Inches(2.0), Inches(5.6), Inches(1.3),
        ['cd "D:\\Web Service\\budgetagent"',
         "python -m venv venv",
         ".\\venv\\Scripts\\pip install --no-index",
         "  --find-links=wheels -r requirements.txt"])
txt(s, Inches(7.03), Inches(3.4), Inches(5.6), Inches(0.8),
    [("‧ venv 必須在部署機重建，不可複製！", 12, GREYB, False),
     ("‧ 裝完確認 venv\\Scripts\\waitress-serve.exe 存在", 12, GREYB, False)], space=6)

rect(s, Inches(0.5), Inches(4.5), Inches(12.33), Inches(2.3), PANEL)
rect(s, Inches(0.5), Inches(4.5), Pt(8), Inches(2.3), RED)
txt(s, Inches(0.8), Inches(4.65), Inches(12), Inches(0.4), "⚠ 三個最常見的坑", 15, RED, True)
txt(s, Inches(0.8), Inches(5.15), Inches(12), Inches(1.6),
    [("‧ wheels 是空的：手動建空資料夾沒用，要真的跑 pip download 把它裝滿。", 12.5, GREYB, False),
     ("‧ 版本對不上：venv 是 3.14、wheels 卻是 cp313 → No matching distribution。先 python --version 查 venv。", 12.5, GREYB, False),
     ("‧ waitress-serve.exe 不存在：多半是 requirements.txt 漏寫 waitress，補上重裝。", 12.5, GREYB, False)], space=10)

# =================================================================
# 7 Step 3  ⭐ 大坑
# =================================================================
s = slide()
title_bar(s, "STEP 3 ‧ 設定檔", "wsgi.py + web.config（最大踩坑點）", "07")

panel(s, Inches(0.5), Inches(1.4), Inches(6.1), Inches(2.3), "backend\\wsgi.py（純內容）", CYAN)
codebox(s, Inches(0.7), Inches(2.0), Inches(5.7), Inches(1.5),
        ["import os, sys",
         "sys.path.insert(0,",
         "  os.path.dirname(",
         "    os.path.abspath(__file__)))",
         "from app import app as application"])

panel(s, Inches(6.83), Inches(1.4), Inches(6.0), Inches(2.3), "web.config 關鍵欄位", GOLD)
txt(s, Inches(7.03), Inches(2.0), Inches(5.6), Inches(1.6),
    [("‧ processPath = venv 內 waitress-serve.exe 完整路徑", 12, GREYB, False),
     ("‧ PYTHONPATH = app.py 所在資料夾（本例 backend）", 12, GREYB, False),
     ("‧ forwardWindowsAuthToken=\"true\"（SSO 必須）", 12, GREYB, False),
     ("‧ 別放 <authentication> 區段（會被鎖→改用 IIS 畫面設）", 12, GREYB, False)], space=8)

rect(s, Inches(0.5), Inches(3.95), Inches(12.33), Inches(2.85), PANEL)
rect(s, Inches(0.5), Inches(3.95), Pt(8), Inches(2.85), RED)
txt(s, Inches(0.8), Inches(4.1), Inches(12), Inches(0.45), "⚠⚠ 血淚教訓：建檔只放「純內容」", 16, RED, True)
txt(s, Inches(0.8), Inches(4.65), Inches(11.9), Inches(2.0),
    [("用記事本建檔，貼上檔案內容本身就好。", 13.5, WHITE, True),
     ("不要把 PowerShell 的包裝指令（@'  ...  '@ | Set-Content）一起存進檔案！", 13, GOLD, False),
     ("否則 web.config 開頭變成 @' 而不是 <?xml → IIS 解析失敗 → 500、且 log 不生成。", 12.5, GREYB, False),
     ("存完用 type web.config 確認：開頭是 <?xml、結尾是 </configuration>、沒有 @'。", 12.5, GREEN, False),
     ("web.config 放在專案最外層（跟 venv / wheels 同一層）。", 12.5, GREYB, False)], space=9)

# =================================================================
# 8 Step 4
# =================================================================
s = slide()
title_bar(s, "STEP 4 ‧ IIS 設定", "建立 AppPool + 網站 + 驗證", "08")

panel(s, Inches(0.5), Inches(1.4), Inches(4.0), Inches(3.0), "① AppPool（引擎）", BLUE)
txt(s, Inches(0.7), Inches(2.0), Inches(3.6), Inches(2.3),
    [("名稱 Pool-budgetagent", 12.5, WHITE, True),
     ("", 5, MUTE, False),
     (".NET CLR 版本：", 12, GREYB, False),
     ("沒有 Managed 程式碼", 12.5, GOLD, True),
     ("", 5, MUTE, False),
     ("啟動模式：", 12, GREYB, False),
     ("AlwaysRunning", 12.5, GOLD, True)], space=6)

panel(s, Inches(4.66), Inches(1.4), Inches(4.0), Inches(3.0), "② 網站", CYAN)
txt(s, Inches(4.86), Inches(2.0), Inches(3.6), Inches(2.3),
    [("名稱 budgetagent", 12.5, WHITE, True),
     ("", 5, MUTE, False),
     ("實體路徑：", 12, GREYB, False),
     ("D:\\Web Service\\budgetagent", 11, GOLD, True),
     ("（最外層）", 10.5, MUTE, False),
     ("", 5, MUTE, False),
     ("應用程式集區：Pool-budgetagent", 11.5, GREYB, False),
     ("Port：8001", 12.5, WHITE, True)], space=5)

panel(s, Inches(8.83), Inches(1.4), Inches(4.0), Inches(3.0), "③ 驗證", GREEN)
txt(s, Inches(9.03), Inches(2.0), Inches(3.6), Inches(2.3),
    [("網站 → 雙擊「驗證」", 12.5, WHITE, True),
     ("（中間切到「功能檢視」", 10.5, MUTE, False),
     ("  才看得到圖示）", 10.5, MUTE, False),
     ("", 6, MUTE, False),
     ("匿名驗證 → 停用", 13, RED, True),
     ("Windows 驗證 → 啟用", 13, GREEN, True),
     ("", 5, MUTE, False),
     ("提供者含 Negotiate/NTLM", 11, GREYB, False)], space=5)

rect(s, Inches(0.5), Inches(4.6), Inches(12.33), Inches(2.2), PANEL)
rect(s, Inches(0.5), Inches(4.6), Pt(8), Inches(2.2), GOLD)
txt(s, Inches(0.8), Inches(4.75), Inches(12), Inches(0.4), "💡 小提醒", 15, GOLD, True)
txt(s, Inches(0.8), Inches(5.2), Inches(11.9), Inches(1.5),
    [("‧ 找不到「網站」項目？在左欄展開「電腦名稱」，它就在「應用程式集區」下方。", 12.5, GREYB, False),
     ("‧ 看不到「驗證」圖示？中間下方從「內容檢視」切到「功能檢視」。", 12.5, GREYB, False),
     ("‧ AppPool ≠ 網站：要分別建立，兩個都要有。", 12.5, GREYB, False)], space=10)

# =================================================================
# 9 Step 5 權限
# =================================================================
s = slide()
title_bar(s, "STEP 5 ‧ 權限", "給 AppPool 帳號資料夾權限", "09")

rect(s, Inches(0.5), Inches(1.4), Inches(12.33), Inches(1.7), PANEL)
rect(s, Inches(0.5), Inches(1.4), Pt(8), Inches(1.7), RED)
txt(s, Inches(0.8), Inches(1.55), Inches(12), Inches(0.4), "為什麼要這步？", 15, RED, True)
txt(s, Inches(0.8), Inches(2.0), Inches(11.9), Inches(1.0),
    [("手動跑能開、IIS 跑卻 500 → 因為 IIS 用低權限帳號 IIS AppPool\\Pool-budgetagent，", 13, GREYB, False),
     ("讀不到你的程式 / venv。給它權限就好。UI 設常卡（位置要選本機），用 icacls 最穩。", 13, GREYB, False)], space=8)

panel(s, Inches(0.5), Inches(3.3), Inches(12.33), Inches(2.0), "icacls 一次搞定（系統管理員 PowerShell）", CYAN)
codebox(s, Inches(0.7), Inches(3.9), Inches(11.9), Inches(1.2),
        ['icacls "D:\\Web Service\\budgetagent" /grant "IIS AppPool\\Pool-budgetagent:(OI)(CI)RX" /T',
         'icacls "D:\\Web Service\\budgetagent\\logs" /grant "IIS AppPool\\Pool-budgetagent:(OI)(CI)M" /T'])

txt(s, Inches(0.5), Inches(5.5), Inches(12.33), Inches(1.3),
    [("‧ 主目錄給 RX（讀取+執行）→ IIS 跑得動 Python；logs 給 M（修改）→ 寫得了 log。", 12.5, GREYB, False),
     ("‧ 路徑有空格 → 一定用雙引號包起來。看到「已順利處理 N 個檔案」即成功。", 12.5, GREYB, False),
     ("‧ /T = 連所有子資料夾/檔案一起套用。", 12.5, GREYB, False)], space=9)

# =================================================================
# 10 Step 6 上線
# =================================================================
s = slide()
title_bar(s, "STEP 6 ‧ 上線", "iisreset → 用正確網址測試", "10")

panel(s, Inches(0.5), Inches(1.4), Inches(6.1), Inches(2.6), "用「正確的網址」測（不一定 localhost）", GOLD)
codebox(s, Inches(0.7), Inches(2.0), Inches(5.7), Inches(0.45), ["Get-WebBinding"])
txt(s, Inches(0.7), Inches(2.6), Inches(5.7), Inches(1.3),
    [("‧ 綁 *:8001  → http://localhost:8001 可以", 12.5, GREYB, False),
     ("‧ 綁特定 IP（例 10.10.28.170:8001）→", 12.5, GREYB, False),
     ("   必須用 http://10.10.28.170:8001", 12.5, GOLD, True),
     ("   用 localhost 會 400 Invalid Hostname", 11, MUTE, False)], space=8)

panel(s, Inches(6.83), Inches(1.4), Inches(6.0), Inches(2.6), "讓其他同事連", CYAN)
codebox(s, Inches(7.03), Inches(2.0), Inches(5.6), Inches(0.95),
        ['New-NetFirewallRule -DisplayName "app-8001"',
         "  -Direction Inbound -Protocol TCP",
         "  -LocalPort 8001 -Action Allow"])
txt(s, Inches(7.03), Inches(3.05), Inches(5.6), Inches(0.9),
    [("‧ 同事用 http://機器名稱:8001 連", 12.5, GREYB, False),
     ("‧ 站台加入瀏覽器「近端內部網路」避免彈窗", 12.5, GREYB, False),
     ("‧ SSO 用機器名稱比 IP 準", 12.5, GREYB, False)], space=6)

rect(s, Inches(0.5), Inches(4.3), Inches(12.33), Inches(2.5), PANEL)
rect(s, Inches(0.5), Inches(4.3), Pt(8), Inches(2.5), GREEN)
txt(s, Inches(0.8), Inches(4.45), Inches(12), Inches(0.45), "🎉 出現你的網頁 = 正式上線成功", 16, GREEN, True)
txt(s, Inches(0.8), Inches(5.0), Inches(11.9), Inches(1.7),
    [("這時可以關掉所有 PowerShell 視窗、關 IIS 管理員、甚至登出 Windows，網站都照跑。", 13, GREYB, False),
     ("重開機後也會自動啟動（因為 AppPool 設了 AlwaysRunning）。", 13, GREYB, False),
     ("收尾：測 AD SSO 認得出工號、移除任何 /debug 測試路由。", 13, GREYB, False)], space=10)

# =================================================================
# 11 錯誤速查
# =================================================================
s = slide()
title_bar(s, "TROUBLESHOOT ‧ 速查", "常見錯誤速查表（本次實戰）", "11")

rows = [("407 Proxy authentication", "公司 proxy 擋 pip", "pip download 加 --proxy"),
        ("PowerShell 出現 >>", "引號沒閉合", "Ctrl+C，整行重打、引號成對"),
        ("No matching distribution", "wheels 版本 ≠ venv 版本", "查 venv 版本，重抓對應 wheels"),
        ("waitress-serve.exe 不存在", "requirements 漏 waitress", "補上、重跑 pip install"),
        ("Fatal error in launcher", "venv 從別處複製", "刪 venv，最終路徑重建"),
        ("web.config 父系鎖定", "<authentication> 被鎖", "拿掉該段，改 IIS 畫面設"),
        ("500（log 沒生成）", "web.config 含了 @'/Set-Content", "記事本重存純 XML"),
        ("500（手動跑卻 OK）", "AppPool 帳號沒權限", "icacls 給 RX / M"),
        ("400 Invalid Hostname", "綁特定 IP 卻用 localhost", "改用該 IP，或繫結改 *"),
        ("500.19 (0x80070021)", "模組沒裝 / 區段被鎖", "查模組→重裝MSI / appcmd unlock"),
        ("502 Bad Gateway", "waitress 路徑錯/崩潰", "看 log，檢查 processPath"),
        ("503 Service Unavailable", "AppPool 停了", "IIS 管理員啟動 AppPool")]
y0 = Inches(1.35)
rh = Inches(0.42)
# 表頭
rect(s, Inches(0.5), y0, Inches(12.33), rh, DARK)
for cx, cw, t in [(0.6, 3.5, "現象"), (4.2, 3.6, "原因"), (7.9, 4.8, "解法")]:
    txt(s, Inches(cx), y0, Inches(cw), rh, t, 12.5, CYAN, True, anchor=MSO_ANCHOR.MIDDLE)
for i, (a, b, c) in enumerate(rows):
    y = y0 + rh + i * Inches(0.43)
    if i % 2 == 0:
        rect(s, Inches(0.5), y, Inches(12.33), Inches(0.43), PANEL)
    txt(s, Inches(0.6), y, Inches(3.5), Inches(0.43), a, 11, WHITE, False, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, Inches(4.2), y, Inches(3.6), Inches(0.43), b, 11, GREYB, False, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, Inches(7.9), y, Inches(4.8), Inches(0.43), c, 11, GREEN, False, anchor=MSO_ANCHOR.MIDDLE)

# =================================================================
# 12 黃金除錯流程
# =================================================================
s = slide()
title_bar(s, "DEBUG ‧ 黃金流程", "卡關時，照這四步找真兇", "12")

debug = [("1", "先手動跑，繞過 IIS", "看 Python 真正的錯誤",
          'python.exe -c "import wsgi; print(\'OK\')"', BLUE),
         ("2", "模擬 IIS 環境啟動", "能開→純 IIS 問題",
          "waitress-serve.exe --port=8003 wsgi:application", CYAN),
         ("3", "看 IIS 啟動失敗紀錄", "事件檢視器找原因",
          "Get-WinEvent -LogName Application ...", GOLD),
         ("4", "看程式 log", "最後的真相",
          'type "...\\logs\\python.log"', GREEN)]
for i, (n, h, d, code, c) in enumerate(debug):
    y = Inches(1.5) + i * Inches(1.05)
    rect(s, Inches(0.5), y, Inches(12.33), Inches(0.92), PANEL)
    rect(s, Inches(0.5), y, Pt(8), Inches(0.92), c)
    rect(s, Inches(0.75), y + Inches(0.21), Inches(0.5), Inches(0.5), c, shape=MSO_SHAPE.OVAL)
    txt(s, Inches(0.75), y + Inches(0.21), Inches(0.5), Inches(0.5), n, 18, BG, True,
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, Inches(1.5), y + Inches(0.12), Inches(4.0), Inches(0.4), h, 14, WHITE, True)
    txt(s, Inches(1.5), y + Inches(0.5), Inches(4.0), Inches(0.35), d, 11, MUTE, False)
    rect(s, Inches(5.6), y + Inches(0.22), Inches(7.0), Inches(0.5), RGBColor(0x0C,0x0E,0x12), line=DARK)
    txt(s, Inches(5.75), y + Inches(0.22), Inches(6.8), Inches(0.5),
        [(code, 11.5, CODE, False, MONO)], anchor=MSO_ANCHOR.MIDDLE)

txt(s, Inches(0.5), Inches(5.95), Inches(12.33), Inches(1.2),
    [("速記版流程：", 13, CYAN, True),
     ("有網路電腦 pip download 裝滿 wheels → 部署機建 venv 離線裝（含 waitress）→ 寫 wsgi/web.config（純內容）", 11.5, GREYB, False),
     ("→ IIS 建 AppPool+網站+驗證 → icacls 授權 → iisreset → 用正確 IP/主機名開網站 🎉", 11.5, GREYB, False)], space=6)

# =================================================================
out = "Flask_IIS部署新手操作簡報.pptx"
prs.save(out)
print("saved:", out, "slides:", len(prs.slides._sldIdLst))
