# -*- coding: utf-8 -*-
"""進出水平衡計算對照系統 — 黑底專案進度報告（總覽 + 10-step 里程碑 + 系統架構流程圖）。"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR

# ---- 黑底主題色 ----
BG    = RGBColor(0x10, 0x12, 0x16)   # 近黑底
PANEL = RGBColor(0x1B, 0x1F, 0x27)   # 面板深灰
DARK  = RGBColor(0x2A, 0x2F, 0x39)   # STEP / phase 深灰
GREYB = RGBColor(0xC9, 0xCF, 0xD8)   # 標題淺灰條
GREYD = RGBColor(0xB2, 0xB9, 0xC4)   # 說明淺灰條
CYAN  = RGBColor(0x35, 0xD0, 0xE6)   # 強調青
BLUE  = RGBColor(0x3E, 0x8F, 0xD6)
GOLD  = RGBColor(0xE7, 0xB4, 0x3C)
GREEN = RGBColor(0x4F, 0xC1, 0x7A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK   = RGBColor(0x18, 0x1C, 0x22)   # 淺條上的深字
MUTE  = RGBColor(0x8A, 0x93, 0xA1)   # 次要灰字
FONT  = "Microsoft JhengHei"

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def slide():
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    r.fill.solid(); r.fill.fore_color.rgb = BG; r.line.fill.background()
    r.shadow.inherit = False
    return s


def rect(s, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE, line=None, lw=1.0):
    sp = s.shapes.add_shape(shape, x, y, w, h)
    sp.fill.solid(); sp.fill.fore_color.rgb = color
    if line: sp.line.color.rgb = line; sp.line.width = Pt(lw)
    else: sp.line.fill.background()
    sp.shadow.inherit = False
    return sp


def txt(s, x, y, w, h, lines, size=14, color=WHITE, bold=False,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space=3):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_top = 0; tf.margin_bottom = 0; tf.margin_left = Pt(2); tf.margin_right = Pt(2)
    if isinstance(lines, str): lines = [(lines, size, color, bold)]
    for i, (t, sz, c, b) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(space); p.space_before = Pt(0)
        r = p.add_run(); r.text = t
        r.font.size = Pt(sz); r.font.bold = b; r.font.color.rgb = c; r.font.name = FONT
    return tb


def title_bar(s, kicker, title, page):
    rect(s, 0, 0, Inches(0.18), Inches(1.05), CYAN)
    txt(s, Inches(0.5), Inches(0.14), Inches(10), Inches(0.32), kicker, 12, CYAN, True)
    txt(s, Inches(0.5), Inches(0.42), Inches(11), Inches(0.6), title, 26, WHITE, True)
    txt(s, Inches(11.8), Inches(0.3), Inches(1.2), Inches(0.5), f"PAGE {page}", 12, MUTE,
        True, align=PP_ALIGN.RIGHT)
    rect(s, Inches(0.5), Inches(1.08), Inches(12.33), Pt(2), DARK)


# =================================================================
# 投影片 1：專案總覽 + 開發團隊
# =================================================================
s = slide()
title_bar(s, "PROJECT PROGRESS REPORT ‧ 2026/06", "進出水平衡計算對照系統", "01")

# 副標
txt(s, Inches(0.5), Inches(1.25), Inches(12.3), Inches(0.4),
    "精確計算蒸發量 ‧ Agent 自動化平衡計算與申報整理　|　預計上線 2026/10　|　目前進度約 25%",
    13, CYAN)

# 左：專案目的
rect(s, Inches(0.5), Inches(1.85), Inches(6.0), Inches(2.0), PANEL)
rect(s, Inches(0.5), Inches(1.85), Inches(6.0), Inches(0.42), DARK)
txt(s, Inches(0.7), Inches(1.85), Inches(5.6), Inches(0.42), "專案目的", 15, CYAN, True,
    anchor=MSO_ANCHOR.MIDDLE)
txt(s, Inches(0.7), Inches(2.42), Inches(5.6), Inches(1.35),
    [("‧ 水平衡計算：彙整進／出水資料，建立可稽核對照基準。", 12.5, GREYB, False),
     ("‧ 蒸發量推算：以平衡差值精準反推蒸發量，取代估算。", 12.5, GREYB, False),
     ("‧ Agent 自動化：代理執行計算並整理後續申報資料。", 12.5, GREYB, False)], space=7)

# 右：專案效益
rect(s, Inches(6.83), Inches(1.85), Inches(6.0), Inches(2.0), PANEL)
rect(s, Inches(6.83), Inches(1.85), Inches(6.0), Inches(0.42), DARK)
txt(s, Inches(7.03), Inches(1.85), Inches(5.6), Inches(0.42), "專案效益", 15, CYAN, True,
    anchor=MSO_ANCHOR.MIDDLE)
bens = [("自動化", "減少人工試算與資料搬運"), ("準確性", "邏輯統一、結果可重現可稽核"),
        ("效率提升", "申報資料自動彙整成型"), ("可追溯／合規", "保留計算過程供申報佐證")]
for i, (h, d) in enumerate(bens):
    r, c = divmod(i, 2)
    x = Inches(7.03) + c * Inches(2.95); y = Inches(2.45) + r * Inches(0.66)
    rect(s, x, y, Pt(7), Inches(0.55), CYAN)
    txt(s, x + Inches(0.15), y, Inches(2.7), Inches(0.55),
        [(h, 12.5, WHITE, True), (d, 10, MUTE, False)], space=1, anchor=MSO_ANCHOR.MIDDLE)

# 當前進度
rect(s, Inches(0.5), Inches(4.0), Inches(12.33), Inches(0.95), PANEL)
txt(s, Inches(0.7), Inches(4.08), Inches(4), Inches(0.3), "當前進度", 14, CYAN, True)
rect(s, Inches(0.7), Inches(4.45), Inches(6.0), Inches(0.2), DARK)
rect(s, Inches(0.7), Inches(4.45), Inches(1.5), Inches(0.2), CYAN)
txt(s, Inches(6.8), Inches(4.4), Inches(1), Inches(0.3), "約 25%", 11, CYAN, True)
txt(s, Inches(8.0), Inches(4.08), Inches(4.7), Inches(0.85),
    [("‧ AS-IS / TO-BE 流程圖整理（進行中）", 11.5, GREYB, False),
     ("‧ 過往相關資料盤點與蒐集（進行中）", 11.5, GREYB, False)], space=4)

# 開發團隊
rect(s, Inches(0.5), Inches(5.1), Inches(12.33), Inches(1.95), PANEL)
rect(s, Inches(0.5), Inches(5.1), Inches(12.33), Inches(0.42), DARK)
txt(s, Inches(0.7), Inches(5.1), Inches(8), Inches(0.42), "開發團隊", 15, CYAN, True,
    anchor=MSO_ANCHOR.MIDDLE)
txt(s, Inches(8.5), Inches(5.1), Inches(4.1), Inches(0.42),
    "＊姓名待補，先以角色呈現", 10.5, MUTE, False, align=PP_ALIGN.RIGHT,
    anchor=MSO_ANCHOR.MIDDLE)
team = [("專案經理 PM", "統籌規劃、時程控管與利害關係人溝通", CYAN),
        ("AI / Agent 工程師", "Agent 設計與平衡計算邏輯實作", BLUE),
        ("資料工程師", "資料盤點、清洗與跨系統串接", GREEN),
        ("領域 / 水務專家", "水平衡與蒸發量計算驗證、申報需求", GOLD),
        ("測試 / QA", "準確性驗證與 UAT 使用者驗收", RGBColor(0xC7,0x7D,0xE0))]
tw = Inches(2.38)
for i, (role, duty, c) in enumerate(team):
    x = Inches(0.62) + i * Inches(2.44); y = Inches(5.65)
    rect(s, x, y, tw, Inches(1.25), BG, line=DARK)
    rect(s, x, y, tw, Pt(5), c)
    rect(s, x + Inches(0.18), y + Inches(0.18), Inches(0.5), Inches(0.5),
         DARK, shape=MSO_SHAPE.OVAL)
    txt(s, x + Inches(0.18), y + Inches(0.18), Inches(0.5), Inches(0.5),
        "▣", 16, c, True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, x + Inches(0.78), y + Inches(0.18), tw - Inches(0.85), Inches(0.55),
        role, 12.5, WHITE, True, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, x + Inches(0.18), y + Inches(0.74), tw - Inches(0.34), Inches(0.45),
        duty, 9.5, MUTE, False)

txt(s, Inches(0.5), Inches(7.18), Inches(12.33), Inches(0.25),
    "下一步：凍結 TO-BE 需求範圍 ▶ 定義平衡計算公式與 Agent 規格 ▶ 完成歷史資料盤點（需相關單位協助提供資料來源與申報欄位）",
    10, MUTE, True, align=PP_ALIGN.CENTER)

# =================================================================
# 投影片 2：里程碑 10-step（仿提供樣式，黑底）
# =================================================================
s = slide()
title_bar(s, "MILESTONE ROADMAP ‧ 2026/06 → 2026/10", "專案里程碑（10 步驟）", "02")
txt(s, Inches(11.0), Inches(1.2), Inches(1.83), Inches(0.3), "完成後應產出", 12, CYAN, True,
    align=PP_ALIGN.RIGHT)

# 階段定義
phases = [
    ("需求確認", "6 月", CYAN, [
        ("01", "10%", "Agent 任務確認", "明確定義水平衡／蒸發量計算範圍、痛點與 Agent 效益"),
        ("02", "20%", "作業流程定義", "將現行人工水量試算與申報流程，翻譯成 Agent 執行節點（AS-IS→TO-BE）")]),
    ("知識建立", "7 月", BLUE, [
        ("03", "30%", "資料文件盤點", "盤點歷史進出水量、用水紀錄與既有申報報表等所需資料"),
        ("04", "40%", "資料整理 / 清洗", "整理相關資料、去除雜訊、統一欄位格式，建立計算基準"),
        ("05", "50%", "結構化分類 / 屬性", "建立水量項目分類與屬性標籤，設定計算與檢索邊界")]),
    ("計算邏輯與串接", "8 月", GOLD, [
        ("06", "60%", "計算邏輯 / Prompt 調校", "定義平衡計算公式與蒸發量推算邏輯，反覆調校降低判讀錯誤"),
        ("07", "70%", "跨系統串接測試", "串接資料來源與申報系統，賦予 Agent 取數與輸出的「手腳」")]),
    ("驗證與上線", "9–10 月", GREEN, [
        ("08", "80%", "準確性驗證", "以歷史實際資料對帳，驗證平衡計算與蒸發量結果無誤"),
        ("09", "90%", "上線測試（UAT）", "平行測試，確認結果符合 Agent 定義且 UX 符合預期"),
        ("10", "100%", "正式上線回饋追蹤", "正式串接、公告上線，產出申報資料並持續追蹤優化")])]

top = Inches(1.55); row_h = Inches(0.5); gap_r = Inches(0.07); gap_p = Inches(0.12)
px = Inches(0.5)                 # phase 欄
sx = Inches(1.55)                # step 欄
tx = Inches(2.75)                # 標題欄
dx = Inches(5.05)                # 說明欄
d_end = Inches(11.6)
y = top
bracket_x = Inches(11.75)
for pname, month, pc, steps in phases:
    block_h = Emu(len(steps) * int(row_h) + (len(steps) - 1) * int(gap_r))
    # phase 欄
    rect(s, px, y, Inches(0.95), block_h, DARK)
    rect(s, px, y, Pt(5), block_h, pc)
    txt(s, px, y, Inches(0.95), block_h,
        [(pname, 13, WHITE, True), (month, 10, pc, True)], space=2,
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    ry = y
    for num, pct, t, d in steps:
        # step box
        rect(s, sx, ry, Inches(1.05), row_h, PANEL, line=DARK)
        txt(s, sx, ry + Inches(0.02), Inches(0.78), row_h,
            [("STEP", 8, MUTE, True), (num, 16, WHITE, True)], space=0,
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        txt(s, sx + Inches(0.74), ry, Inches(0.3), row_h, pct, 9.5, pc, True,
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        # 標題條
        rect(s, tx, ry, Inches(2.2), row_h, GREYB)
        txt(s, tx + Inches(0.12), ry, Inches(2.0), row_h, t, 12.5, INK, True,
            anchor=MSO_ANCHOR.MIDDLE)
        # 說明條
        rect(s, dx, ry, Emu(int(d_end) - int(dx)), row_h, GREYD)
        txt(s, dx + Inches(0.15), ry, Emu(int(d_end) - int(dx)) - Inches(0.3), row_h,
            d, 11, INK, False, anchor=MSO_ANCHOR.MIDDLE)
        ry = Emu(int(ry) + int(row_h) + int(gap_r))
    # 右側 bracket → 完成後產出
    rect(s, bracket_x, y, Pt(3), block_h, pc)
    y = Emu(int(y) + int(block_h) + int(gap_p))

# 完成後應產出（對應四階段）
outs = ["需求規格 / AS-IS·TO-BE 流程圖", "知識庫 / 結構化資料集",
        "平衡計算引擎 + Agent 串接", "驗證報告 / 上線版申報資料"]
oy = top
for (pname, month, pc, steps), o in zip(phases, outs):
    block_h = Emu(len(steps) * int(row_h) + (len(steps) - 1) * int(gap_r))
    txt(s, Inches(11.95), oy, Inches(1.35), block_h, o, 10, pc, True,
        anchor=MSO_ANCHOR.MIDDLE)
    oy = Emu(int(oy) + int(block_h) + int(gap_p))

# =================================================================
# 投影片 3：系統架構流程圖
# =================================================================
s = slide()
title_bar(s, "SYSTEM ARCHITECTURE FLOW", "系統架構流程圖", "03")


def fbox(s, x, y, w, h, head, sub, c):
    rect(s, x, y, w, h, PANEL, line=c, lw=1.5)
    rect(s, x, y, w, Pt(5), c)
    txt(s, x + Inches(0.12), y + Inches(0.1), w - Inches(0.24), h - Inches(0.2),
        [(head, 12.5, WHITE, True), (sub, 9.5, MUTE, False)], space=3,
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def arrow(s, x, y, w, color=CYAN):
    a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x, y, w, Inches(0.32))
    a.fill.solid(); a.fill.fore_color.rgb = color; a.line.fill.background()
    a.shadow.inherit = False


# 主管線：5 個方塊
bw, bh, by = Inches(2.2), Inches(1.25), Inches(2.9)
xs = [Inches(0.55), Inches(3.15), Inches(5.75), Inches(8.35), Inches(10.95)]
main = [("① 資料來源", "SCADA / 水錶 /\n進出水量 / 歷史報表", BLUE),
        ("② 資料整理 + 知識庫", "清洗統一格式\nRAG 結構化分類", GREEN),
        ("③ Agent 計算引擎", "進出水平衡計算\n蒸發量推算", CYAN),
        ("④ 準確性驗證", "歷史資料對帳\n異常檢核", GOLD),
        ("⑤ 申報資料整理", "自動彙整\n產出申報報表", RGBColor(0xC7,0x7D,0xE0))]
for x, (h, sub, c) in zip(xs, main):
    fbox(s, x, by, bw, bh, h, sub.replace("\n", "　"), c)
for i in range(4):
    ax = Emu(int(xs[i]) + int(bw) + Inches(0.04))
    arrow(s, ax, Emu(int(by) + int(bh)//2 - int(Inches(0.16))), Inches(0.52))

# Agent 代理層（上方橫帶，涵蓋計算引擎）
rect(s, Inches(5.6), Inches(1.85), Inches(2.5), Inches(0.55), DARK, line=CYAN)
txt(s, Inches(5.6), Inches(1.85), Inches(2.5), Inches(0.55),
    "Agent 代理層（自動編排）", 11, CYAN, True, align=PP_ALIGN.CENTER,
    anchor=MSO_ANCHOR.MIDDLE)
conn = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(6.85), Inches(2.4),
                              Inches(6.85), Inches(2.9))
conn.line.color.rgb = CYAN; conn.line.width = Pt(1.25); conn.line.dash_style = 2

# 輸出
fbox(s, Inches(10.95), Inches(4.85), bw, bh, "⑥ 輸出 / 交付",
     "申報報表　管理儀表板", GREEN)
ca = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(12.05), Inches(4.15),
                            Inches(12.05), Inches(4.85))
ca.line.color.rgb = GREEN; ca.line.width = Pt(2)

# 回饋迴路（驗證 → 計算引擎）
fb = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(9.45), Inches(4.4),
                            Inches(6.85), Inches(4.4))
fb.line.color.rgb = GOLD; fb.line.width = Pt(1.5); fb.line.dash_style = 2
fb2 = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(9.45), Inches(4.15),
                             Inches(9.45), Inches(4.4))
fb2.line.color.rgb = GOLD; fb2.line.width = Pt(1.5); fb2.line.dash_style = 2
fb3 = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(6.85), Inches(4.4),
                             Inches(6.85), Inches(4.15))
fb3.line.color.rgb = GOLD; fb3.line.width = Pt(1.5); fb3.line.dash_style = 2
txt(s, Inches(7.0), Inches(4.45), Inches(2.4), Inches(0.3),
    "驗證異常 → 回饋修正", 9.5, GOLD, True, align=PP_ALIGN.CENTER)

# 底部圖例 / 說明
rect(s, Inches(0.55), Inches(5.6), Inches(9.9), Inches(1.4), PANEL)
txt(s, Inches(0.75), Inches(5.7), Inches(9.5), Inches(0.3), "流程說明", 13, CYAN, True)
txt(s, Inches(0.75), Inches(6.08), Inches(9.5), Inches(0.85),
    [("‧ 資料自來源彙整後經清洗與結構化進入知識庫，作為計算依據。", 11, GREYB, False),
     ("‧ Agent 代理層自動編排平衡計算與蒸發量推算，結果經準確性驗證；異常則回饋修正後重算。", 11, GREYB, False),
     ("‧ 驗證通過之數據自動整理為申報資料，輸出報表與管理儀表板。", 11, GREYB, False)], space=5)

txt(s, Inches(0.55), Inches(7.15), Inches(12.3), Inches(0.25),
    "實線＝主資料流　|　虛線＝回饋／代理控制流", 9.5, MUTE, False, align=PP_ALIGN.CENTER)

prs.save("進出水平衡計算對照系統_專案進度報告.pptx")
print("OK", len(prs.slides._sldIdLst), "slides")
